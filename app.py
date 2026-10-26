from __future__ import annotations

import csv
import io
import json
import queue
import re
import sqlite3
import threading
import time
import webbrowser
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Any

from flask import Flask, flash, jsonify, redirect, render_template, request, url_for
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "prospect.db"
CSV_PATH = BASE_DIR / "clientes.csv"
HISTORY_PATH = BASE_DIR / "historico.json"
PROFILE_PATH = BASE_DIR / "perfil_whatsapp"
DEFAULT_MESSAGE = (
    "Olá, {contato}, tudo bem?\n\n"
    "Meu nome é Alex e trabalho com desenvolvimento de sites e sistemas "
    "personalizados para empresas, criando soluções digitais de acordo "
    "com as necessidades de cada negócio, com foco em organização, "
    "eficiência e otimização de processos.\n\n"
    "Gostaria de me colocar à disposição para conversar e entender se "
    "alguma solução digital pode agregar valor à {empresa}.\n\n"
    "Fico à disposição!\n\nAtenciosamente,\nAlex\n\n"
    "Se não quiser receber novas mensagens, é só me avisar."
)
CONSENT_YES = {"sim", "s", "yes", "true", "1"}
app = Flask(__name__)
app.secret_key = "local-dashboard-change-this-key"


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def consent_yes(value: Any) -> bool:
    return str(value or "").strip().lower() in CONSENT_YES


@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db() -> None:
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            company TEXT NOT NULL,
            contact_name TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL UNIQUE,
            consent INTEGER NOT NULL DEFAULT 0,
            blocked INTEGER NOT NULL DEFAULT 0,
            notes TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS campaigns (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            template_id INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY(template_id) REFERENCES templates(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS campaign_contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            contact_id INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            prepared_message TEXT NOT NULL DEFAULT '',
            prepared_at TEXT,
            sent_at TEXT,
            error TEXT NOT NULL DEFAULT '',
            UNIQUE(campaign_id, contact_id),
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE,
            FOREIGN KEY(contact_id) REFERENCES contacts(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS templates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            body TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS activity (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER,
            contact_id INTEGER,
            event TEXT NOT NULL,
            detail TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE SET NULL,
            FOREIGN KEY(contact_id) REFERENCES contacts(id) ON DELETE SET NULL
        );
        CREATE INDEX IF NOT EXISTS idx_campaign_contacts_status ON campaign_contacts(campaign_id, status);
        CREATE INDEX IF NOT EXISTS idx_activity_created ON activity(created_at);
        """)
        if con.execute("SELECT COUNT(*) FROM templates").fetchone()[0] == 0:
            con.execute("INSERT INTO templates(name, body, created_at, updated_at) VALUES (?, ?, ?, ?)",
                        ("Prospecção inicial", DEFAULT_MESSAGE, now(), now()))


def log_event(event: str, detail: str = "", campaign_id: int | None = None, contact_id: int | None = None) -> None:
    with db() as con:
        con.execute("INSERT INTO activity(campaign_id, contact_id, event, detail, created_at) VALUES (?, ?, ?, ?, ?)",
                    (campaign_id, contact_id, event, detail, now()))


def migrate_legacy_files() -> None:
    """Import the original CSV and JSON history once, without overwriting current records."""
    with db() as con:
        imported = con.execute("SELECT value FROM app_meta WHERE key='legacy_migrated'").fetchone() if table_exists(con, "app_meta") else None
        if not table_exists(con, "app_meta"):
            con.execute("CREATE TABLE IF NOT EXISTS app_meta(key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        if imported:
            return
        if CSV_PATH.exists():
            try:
                with CSV_PATH.open("r", newline="", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f)
                    expected = {"empresa", "contato", "telefone", "autorizou"}
                    if reader.fieldnames and expected.issubset(set(reader.fieldnames)):
                        for row in reader:
                            company = (row.get("empresa") or "").strip()
                            phone = digits(row.get("telefone"))
                            if not company or not 10 <= len(phone) <= 15:
                                continue
                            con.execute("""INSERT OR IGNORE INTO contacts(company, contact_name, phone, consent, created_at, updated_at)
                                          VALUES (?, ?, ?, ?, ?, ?)""",
                                        (company, (row.get("contato") or "").strip(), phone,
                                         int(consent_yes(row.get("autorizou"))), now(), now()))
            except (OSError, csv.Error):
                pass
        # Keep the legacy JSON available as an archive. Old statuses are imported only as activity,
        # never as proof of a successful send.
        if HISTORY_PATH.exists():
            try:
                legacy = json.loads(HISTORY_PATH.read_text(encoding="utf-8"))
                for key, item in legacy.items() if isinstance(legacy, dict) else []:
                    phone = digits(item.get("telefone")) if isinstance(item, dict) else ""
                    old_status = item.get("status", "registro legado") if isinstance(item, dict) else "registro legado"
                    campaign_name = str(key).split(":", 1)[0].strip() or "Campanha legada"
                    if not phone:
                        continue
                    contact = con.execute("SELECT id FROM contacts WHERE phone=?", (phone,)).fetchone()
                    if not contact:
                        continue
                    campaign = con.execute("SELECT id FROM campaigns WHERE name=? ORDER BY id LIMIT 1", (campaign_name,)).fetchone()
                    if not campaign:
                        template = con.execute("SELECT id FROM templates ORDER BY id LIMIT 1").fetchone()
                        cur = con.execute("INSERT INTO campaigns(name, template_id, created_at) VALUES (?, ?, ?)",
                                          (campaign_name, template["id"] if template else None, now()))
                        campaign_id = cur.lastrowid
                    else:
                        campaign_id = campaign["id"]
                    con.execute("""INSERT OR IGNORE INTO campaign_contacts(campaign_id, contact_id, status, error)
                                  VALUES (?, ?, 'legacy_review', ?)""",
                                (campaign_id, contact["id"], f"Histórico antigo ({old_status}): envio não confirmado. Revise antes de continuar."))
                    con.execute("INSERT INTO activity(campaign_id, contact_id, event, detail, created_at) VALUES (?, ?, ?, ?, ?)",
                                (campaign_id, contact["id"], "legacy_history", f"Histórico migrado ({old_status}); envio não confirmado", now()))
            except (OSError, json.JSONDecodeError):
                pass
        con.execute("INSERT OR REPLACE INTO app_meta(key, value) VALUES ('legacy_migrated', '1')")


def table_exists(con: sqlite3.Connection, name: str) -> bool:
    return con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,)).fetchone() is not None


def render_message(body: str, contact: sqlite3.Row | dict) -> str:
    values = {
        "empresa": contact["company"],
        "contato": contact["contact_name"] or "",
        "telefone": contact["phone"],
    }
    class Safe(dict):
        def __missing__(self, key):
            return "{" + key + "}"
    return body.format_map(Safe(values)).replace("Olá, , tudo bem?", "Olá, tudo bem?")


class WhatsAppWorker:
    """Owns Playwright in a single worker thread; never clicks Send automatically."""
    def __init__(self):
        self.commands: queue.Queue[tuple[str, dict]] = queue.Queue()
        self.lock = threading.Lock()
        self.state = {"connected": False, "busy": False, "message": "WhatsApp desconectado", "last_error": ""}
        self.thread = threading.Thread(target=self._run, daemon=True, name="whatsapp-worker")
        self.thread.start()

    def snapshot(self) -> dict:
        with self.lock:
            return dict(self.state)

    def set_state(self, **kwargs):
        with self.lock:
            self.state.update(kwargs)

    def command(self, name: str, **data):
        self.commands.put((name, data))

    def _run(self):
        pw = context = page = None
        while True:
            name, data = self.commands.get()
            try:
                if name == "connect":
                    if context:
                        try:
                            context.close()
                        except Exception:
                            pass
                    pw = sync_playwright().start()
                    context = pw.chromium.launch_persistent_context(
                        user_data_dir=str(PROFILE_PATH), headless=False,
                        viewport={"width": 1280, "height": 900}, args=["--start-maximized"])
                    page = context.pages[0] if context.pages else context.new_page()
                    page.goto("https://web.whatsapp.com/", wait_until="domcontentloaded")
                    self.set_state(connected=False, busy=True, message="Aguardando login no WhatsApp Web…", last_error="")
                    try:
                        page.get_by_role("textbox", name=re.compile("Pesquisar|Search", re.I)).wait_for(timeout=120_000)
                        self.set_state(connected=True, busy=False, message="WhatsApp Web conectado. O envio continua manual.")
                    except PlaywrightTimeoutError:
                        self.set_state(connected=False, busy=False, message="Login não confirmado. Verifique o QR Code no navegador.")
                elif name == "prepare":
                    if not context or not page:
                        self.set_state(message="Conecte o WhatsApp Web antes de preparar uma mensagem.", last_error="WhatsApp desconectado")
                        continue
                    campaign_contact_id = int(data["campaign_contact_id"])
                    with db() as con:
                        row = con.execute("""SELECT cc.id AS ccid, cc.campaign_id, c.id AS contact_id, c.company,
                            c.contact_name, c.phone, cc.status, cc.prepared_message
                            FROM campaign_contacts cc JOIN contacts c ON c.id=cc.contact_id
                            WHERE cc.id=?""", (campaign_contact_id,)).fetchone()
                        if not row:
                            self.set_state(message="Contato da campanha não encontrado.")
                            continue
                        if row["status"] == "sent":
                            self.set_state(message="Esse contato já está marcado como enviado.")
                            continue
                        if row["status"] == "blocked" or row["status"] == "opted_out":
                            self.set_state(message="Contato bloqueado/descadastrado; não será preparado.")
                            continue
                        if row["status"] == "legacy_review":
                            self.set_state(message="Registro migrado do histórico antigo. Revise e redefina para pendente antes de preparar novamente.")
                            continue
                        if row["status"] == "pending" and con.execute("SELECT consent, blocked FROM contacts WHERE id=?", (row["contact_id"],)).fetchone() != None:
                            ccheck = con.execute("SELECT consent, blocked FROM contacts WHERE id=?", (row["contact_id"],)).fetchone()
                            if not ccheck["consent"] or ccheck["blocked"]:
                                self.set_state(message="Contato sem autorização ou bloqueado; não será preparado.")
                                continue
                    self.set_state(busy=True, message=f"Abrindo conversa com {row['company']}…", last_error="")
                    page.goto(f"https://web.whatsapp.com/send?phone={row['phone']}", wait_until="domcontentloaded")
                    composer = page.locator('footer [contenteditable="true"][role="textbox"]').last
                    composer.wait_for(timeout=45_000)
                    message = row["prepared_message"]
                    if not message:
                        with db() as con:
                            tpl = con.execute("SELECT body FROM templates WHERE id=(SELECT template_id FROM campaigns WHERE id=?)", (row["campaign_id"],)).fetchone()
                            body = tpl["body"] if tpl else DEFAULT_MESSAGE
                            c = con.execute("SELECT company, contact_name, phone FROM contacts WHERE id=?", (row["contact_id"],)).fetchone()
                            message = render_message(body, c)
                            con.execute("UPDATE campaign_contacts SET prepared_message=? WHERE id=?", (message, campaign_contact_id))
                    composer.fill(message)
                    with db() as con:
                        con.execute("UPDATE campaign_contacts SET status='prepared', prepared_at=?, error='' WHERE id=?", (now(), campaign_contact_id))
                        con.execute("INSERT INTO activity(campaign_id, contact_id, event, detail, created_at) VALUES (?, ?, ?, ?, ?)",
                                    (row["campaign_id"], row["contact_id"], "prepared", "Mensagem preenchida; envio ainda não confirmado", now()))
                    self.set_state(connected=True, busy=False, message=f"Mensagem preparada para {row['company']}. Revise e envie manualmente no WhatsApp Web.")
                elif name == "close":
                    if context:
                        context.close()
                    if pw:
                        pw.stop()
                    context = page = pw = None
                    self.set_state(connected=False, busy=False, message="WhatsApp Web desconectado.")
            except Exception as exc:
                self.set_state(busy=False, last_error=str(exc), message=f"Falha na automação: {exc}")
                if name == "connect":
                    self.set_state(connected=False)


worker = WhatsAppWorker()


def stats() -> dict:
    with db() as con:
        return {
            "contacts": con.execute("SELECT COUNT(*) FROM contacts").fetchone()[0],
            "authorized": con.execute("SELECT COUNT(*) FROM contacts WHERE consent=1 AND blocked=0").fetchone()[0],
            "campaigns": con.execute("SELECT COUNT(*) FROM campaigns").fetchone()[0],
            "sent": con.execute("SELECT COUNT(*) FROM campaign_contacts WHERE status='sent'").fetchone()[0],
            "pending": con.execute("SELECT COUNT(*) FROM campaign_contacts WHERE status='pending'").fetchone()[0],
            "prepared": con.execute("SELECT COUNT(*) FROM campaign_contacts WHERE status='prepared'").fetchone()[0],
            "failed": con.execute("SELECT COUNT(*) FROM campaign_contacts WHERE status='failed'").fetchone()[0],
        }


@app.context_processor
def inject_globals():
    return {"wa_state": worker.snapshot(), "app_stats": stats()}


@app.route("/")
def dashboard():
    with db() as con:
        campaigns = con.execute("""SELECT ca.*, COUNT(cc.id) AS total,
            SUM(CASE WHEN cc.status='sent' THEN 1 ELSE 0 END) AS sent,
            SUM(CASE WHEN cc.status='pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN cc.status='prepared' THEN 1 ELSE 0 END) AS prepared
            FROM campaigns ca LEFT JOIN campaign_contacts cc ON cc.campaign_id=ca.id
            GROUP BY ca.id ORDER BY ca.id DESC LIMIT 8""").fetchall()
        activity = con.execute("""SELECT a.*, c.company, c.contact_name FROM activity a
            LEFT JOIN contacts c ON c.id=a.contact_id ORDER BY a.id DESC LIMIT 10""").fetchall()
    return render_template("dashboard.html", campaigns=campaigns, activity=activity)


@app.route("/contacts", methods=["GET", "POST"])
def contacts():
    if request.method == "POST":
        company = request.form.get("company", "").strip()
        name = request.form.get("contact_name", "").strip()
        phone = digits(request.form.get("phone"))
        consent = int(request.form.get("consent") == "on")
        notes = request.form.get("notes", "").strip()
        if not company or not 10 <= len(phone) <= 15:
            flash("Informe a empresa e um telefone válido com código do país e DDD.", "error")
        else:
            try:
                with db() as con:
                    con.execute("INSERT INTO contacts(company, contact_name, phone, consent, notes, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                                (company, name, phone, consent, notes, now(), now()))
                log_event("contact_created", f"Contato cadastrado: {company}")
                flash("Contato cadastrado.", "success")
            except sqlite3.IntegrityError:
                flash("Esse telefone já está cadastrado.", "error")
        return redirect(url_for("contacts"))
    q = request.args.get("q", "").strip()
    with db() as con:
        if q:
            rows = con.execute("SELECT * FROM contacts WHERE company LIKE ? OR contact_name LIKE ? OR phone LIKE ? ORDER BY company", (f"%{q}%", f"%{q}%", f"%{q}%")).fetchall()
        else:
            rows = con.execute("SELECT * FROM contacts ORDER BY id DESC LIMIT 500").fetchall()
    return render_template("contacts.html", contacts=rows, q=q)


@app.post("/contacts/<int:contact_id>/consent")
def toggle_consent(contact_id: int):
    with db() as con:
        row = con.execute("SELECT consent FROM contacts WHERE id=?", (contact_id,)).fetchone()
        if row:
            new = 0 if row["consent"] else 1
            con.execute("UPDATE contacts SET consent=?, updated_at=? WHERE id=?", (new, now(), contact_id))
            log_event("consent_changed", f"Autorização alterada para {new}", contact_id=contact_id)
    return redirect(url_for("contacts"))


@app.post("/contacts/<int:contact_id>/block")
def block_contact(contact_id: int):
    with db() as con:
        con.execute("UPDATE contacts SET blocked=1, consent=0, updated_at=? WHERE id=?", (now(), contact_id))
        con.execute("UPDATE campaign_contacts SET status='opted_out', error='Contato bloqueado/descadastrado' WHERE contact_id=? AND status IN ('pending','prepared')", (contact_id,))
    log_event("contact_blocked", "Contato bloqueado/descadastrado", contact_id=contact_id)
    flash("Contato bloqueado. Não será incluído em novas preparações.", "success")
    return redirect(url_for("contacts"))


def parse_upload(file_storage):
    name = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        values = ws.iter_rows(values_only=True)
        headers = [str(x or "").strip().lower() for x in next(values, [])]
        rows = [dict(zip(headers, [str(v or "").strip() for v in row])) for row in values]
        wb.close()
        return rows
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        return [{str(k or "").strip().lower(): str(v or "").strip() for k, v in row.items()} for row in reader]
    raise ValueError("Formato não suportado. Envie um arquivo CSV ou XLSX.")


@app.post("/contacts/import")
def import_contacts():
    file = request.files.get("file")
    if not file or not file.filename:
        flash("Selecione um CSV ou XLSX.", "error")
        return redirect(url_for("contacts"))
    try:
        rows = parse_upload(file)
        added = updated = skipped = 0
        for row in rows:
            company = (row.get("empresa") or row.get("company") or "").strip()
            contact = (row.get("contato") or row.get("contact") or row.get("contact_name") or "").strip()
            phone = digits(row.get("telefone") or row.get("phone"))
            consent_value = row.get("autorizou", row.get("consent", "nao"))
            consent = int(consent_yes(consent_value))
            if not company or not 10 <= len(phone) <= 15:
                skipped += 1
                continue
            with db() as con:
                exists = con.execute("SELECT id FROM contacts WHERE phone=?", (phone,)).fetchone()
                if exists:
                    con.execute("UPDATE contacts SET company=?, contact_name=?, consent=?, updated_at=? WHERE phone=?",
                                (company, contact, consent, now(), phone))
                    updated += 1
                else:
                    con.execute("INSERT INTO contacts(company, contact_name, phone, consent, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
                                (company, contact, phone, consent, now(), now()))
                    added += 1
        log_event("contacts_imported", f"Importados {added}; atualizados {updated}; ignorados {skipped}")
        flash(f"Importação concluída: {added} novos, {updated} atualizados, {skipped} ignorados.", "success")
    except Exception as exc:
        flash(f"Falha na importação: {exc}", "error")
    return redirect(url_for("contacts"))


@app.route("/templates", methods=["GET", "POST"])
def templates_page():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        body = request.form.get("body", "").strip()
        template_id = request.form.get("template_id", "").strip()
        if not name or not body:
            flash("Informe o nome e o texto do modelo.", "error")
        else:
            with db() as con:
                if template_id:
                    con.execute("UPDATE templates SET name=?, body=?, updated_at=? WHERE id=?", (name, body, now(), int(template_id)))
                else:
                    con.execute("INSERT INTO templates(name, body, created_at, updated_at) VALUES (?, ?, ?, ?)", (name, body, now(), now()))
            flash("Modelo salvo.", "success")
        return redirect(url_for("templates_page"))
    with db() as con:
        templates = con.execute("SELECT * FROM templates ORDER BY id DESC").fetchall()
    return render_template("templates.html", templates=templates)


@app.post("/templates/<int:template_id>/delete")
def delete_template(template_id: int):
    with db() as con:
        used = con.execute("SELECT COUNT(*) FROM campaigns WHERE template_id=?", (template_id,)).fetchone()[0]
        if used:
            flash("Esse modelo está vinculado a campanhas e não pode ser excluído.", "error")
        else:
            con.execute("DELETE FROM templates WHERE id=?", (template_id,))
            flash("Modelo excluído.", "success")
    return redirect(url_for("templates_page"))


@app.route("/campaigns", methods=["GET", "POST"])
def campaigns():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        template_id = request.form.get("template_id", "").strip()
        if not name or not template_id:
            flash("Informe o nome da campanha e selecione um modelo.", "error")
            return redirect(url_for("campaigns"))
        with db() as con:
            cur = con.execute("INSERT INTO campaigns(name, template_id, created_at) VALUES (?, ?, ?)", (name, int(template_id), now()))
            campaign_id = cur.lastrowid
            contacts = con.execute("SELECT id FROM contacts WHERE consent=1 AND blocked=0 ORDER BY company").fetchall()
            for c in contacts:
                con.execute("INSERT OR IGNORE INTO campaign_contacts(campaign_id, contact_id, status) VALUES (?, ?, 'pending')", (campaign_id, c["id"]))
        log_event("campaign_created", f"Campanha criada com {len(contacts)} contatos elegíveis", campaign_id=campaign_id)
        flash(f"Campanha criada com {len(contacts)} contatos autorizados. Confira os contatos antes de iniciar.", "success")
        return redirect(url_for("campaign_detail", campaign_id=campaign_id))
    with db() as con:
        rows = con.execute("""SELECT ca.*, COUNT(cc.id) AS total,
            SUM(CASE WHEN cc.status='sent' THEN 1 ELSE 0 END) AS sent,
            SUM(CASE WHEN cc.status='pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN cc.status='prepared' THEN 1 ELSE 0 END) AS prepared,
            SUM(CASE WHEN cc.status='failed' THEN 1 ELSE 0 END) AS failed
            FROM campaigns ca LEFT JOIN campaign_contacts cc ON cc.campaign_id=ca.id GROUP BY ca.id ORDER BY ca.id DESC""").fetchall()
        templates = con.execute("SELECT id, name FROM templates ORDER BY name").fetchall()
    return render_template("campaigns.html", campaigns=rows, templates=templates)


@app.route("/campaigns/<int:campaign_id>")
def campaign_detail(campaign_id: int):
    with db() as con:
        campaign = con.execute("SELECT ca.*, t.name AS template_name FROM campaigns ca LEFT JOIN templates t ON t.id=ca.template_id WHERE ca.id=?", (campaign_id,)).fetchone()
        if not campaign:
            flash("Campanha não encontrada.", "error")
            return redirect(url_for("campaigns"))
        rows = con.execute("""SELECT cc.*, c.company, c.contact_name, c.phone, c.consent, c.blocked
            FROM campaign_contacts cc JOIN contacts c ON c.id=cc.contact_id WHERE cc.campaign_id=?
            ORDER BY CASE cc.status WHEN 'prepared' THEN 0 WHEN 'pending' THEN 1 WHEN 'failed' THEN 2 ELSE 3 END, c.company""", (campaign_id,)).fetchall()
        templates = con.execute("SELECT id, name FROM templates ORDER BY name").fetchall()
    return render_template("campaign_detail.html", campaign=campaign, rows=rows, templates=templates)


@app.post("/campaigns/<int:campaign_id>/add-eligible")
def add_eligible(campaign_id: int):
    with db() as con:
        contacts = con.execute("SELECT id FROM contacts WHERE consent=1 AND blocked=0").fetchall()
        added = 0
        for c in contacts:
            cur = con.execute("INSERT OR IGNORE INTO campaign_contacts(campaign_id, contact_id, status) VALUES (?, ?, 'pending')", (campaign_id, c["id"]))
            added += cur.rowcount
    log_event("campaign_contacts_added", f"{added} contatos elegíveis adicionados", campaign_id=campaign_id)
    flash(f"{added} novos contatos adicionados à campanha.", "success")
    return redirect(url_for("campaign_detail", campaign_id=campaign_id))


@app.post("/campaign-contacts/<int:cc_id>/prepare")
def prepare_contact(cc_id: int):
    with db() as con:
        row = con.execute("""SELECT cc.*, c.consent, c.blocked FROM campaign_contacts cc JOIN contacts c ON c.id=cc.contact_id WHERE cc.id=?""", (cc_id,)).fetchone()
        if not row:
            flash("Contato não encontrado.", "error")
        elif not row["consent"] or row["blocked"] or row["status"] == "opted_out":
            flash("Contato sem autorização ou bloqueado. Nenhuma mensagem foi preparada.", "error")
        elif row["status"] == "sent":
            flash("Esse contato já está marcado como enviado.", "error")
        elif row["status"] == "legacy_review":
            flash("Esse contato veio do histórico antigo. Revise o envio anterior e use ‘Voltar para pendente’ antes de preparar novamente.", "error")
        else:
            worker.command("prepare", campaign_contact_id=cc_id)
            flash("Preparação solicitada. Confira o estado do WhatsApp no topo e revise a conversa no navegador.", "success")
    return redirect(url_for("campaign_detail", campaign_id=request.form.get("campaign_id", "") or campaign_id_for(cc_id)))


def campaign_id_for(cc_id: int) -> int:
    with db() as con:
        row = con.execute("SELECT campaign_id FROM campaign_contacts WHERE id=?", (cc_id,)).fetchone()
        return int(row[0]) if row else 1


@app.post("/campaign-contacts/<int:cc_id>/confirm-sent")
def confirm_sent(cc_id: int):
    with db() as con:
        row = con.execute("SELECT campaign_id, contact_id, status FROM campaign_contacts WHERE id=?", (cc_id,)).fetchone()
        if row and row["status"] == "prepared":
            con.execute("UPDATE campaign_contacts SET status='sent', sent_at=?, error='' WHERE id=?", (now(), cc_id))
            con.execute("INSERT INTO activity(campaign_id, contact_id, event, detail, created_at) VALUES (?, ?, ?, ?, ?)",
                        (row["campaign_id"], row["contact_id"], "sent_confirmed", "Envio confirmado manualmente pelo usuário", now()))
            flash("Envio registrado como confirmado por você.", "success")
        else:
            flash("Só é possível confirmar um contato cuja mensagem foi preparada.", "error")
    return redirect(url_for("campaign_detail", campaign_id=campaign_id_for(cc_id)))


@app.post("/campaign-contacts/<int:cc_id>/failed")
def mark_failed(cc_id: int):
    error = request.form.get("error", "Envio cancelado ou não confirmado pelo usuário").strip()
    with db() as con:
        row = con.execute("SELECT campaign_id, contact_id FROM campaign_contacts WHERE id=?", (cc_id,)).fetchone()
        if row:
            con.execute("UPDATE campaign_contacts SET status='failed', error=? WHERE id=?", (error, cc_id))
            con.execute("INSERT INTO activity(campaign_id, contact_id, event, detail, created_at) VALUES (?, ?, ?, ?, ?)",
                        (row["campaign_id"], row["contact_id"], "failed", error, now()))
    flash("Contato marcado como não enviado. Você pode tentar novamente depois.", "success")
    return redirect(url_for("campaign_detail", campaign_id=campaign_id_for(cc_id)))


@app.post("/campaign-contacts/<int:cc_id>/reset")
def reset_contact(cc_id: int):
    with db() as con:
        row = con.execute("SELECT campaign_id FROM campaign_contacts WHERE id=?", (cc_id,)).fetchone()
        if row:
            con.execute("UPDATE campaign_contacts SET status='pending', prepared_message='', prepared_at=NULL, sent_at=NULL, error='' WHERE id=? AND status!='sent'", (cc_id,))
    return redirect(url_for("campaign_detail", campaign_id=campaign_id_for(cc_id)))


@app.route("/reports")
def reports():
    with db() as con:
        campaign_rows = con.execute("""SELECT ca.name, ca.created_at, COUNT(cc.id) total,
            SUM(CASE WHEN cc.status='sent' THEN 1 ELSE 0 END) sent,
            SUM(CASE WHEN cc.status='prepared' THEN 1 ELSE 0 END) prepared,
            SUM(CASE WHEN cc.status='pending' THEN 1 ELSE 0 END) pending,
            SUM(CASE WHEN cc.status='failed' THEN 1 ELSE 0 END) failed,
            SUM(CASE WHEN cc.status='opted_out' THEN 1 ELSE 0 END) opted_out
            FROM campaigns ca LEFT JOIN campaign_contacts cc ON cc.campaign_id=ca.id GROUP BY ca.id ORDER BY ca.id DESC""").fetchall()
        recent = con.execute("""SELECT a.*, ca.name AS campaign_name, c.company, c.contact_name FROM activity a
            LEFT JOIN campaigns ca ON ca.id=a.campaign_id LEFT JOIN contacts c ON c.id=a.contact_id ORDER BY a.id DESC LIMIT 100""").fetchall()
    return render_template("reports.html", campaign_rows=campaign_rows, recent=recent)


@app.route("/api/status")
def api_status():
    return jsonify({"whatsapp": worker.snapshot(), "stats": stats()})


@app.post("/api/whatsapp/connect")
def whatsapp_connect():
    worker.command("connect")
    return jsonify({"ok": True, "message": "Abertura do WhatsApp solicitada."})


@app.post("/api/whatsapp/close")
def whatsapp_close():
    worker.command("close")
    return jsonify({"ok": True, "message": "Fechamento solicitado."})


@app.get("/health")
def health():
    return jsonify({"ok": True, "database": str(DB_PATH.name)})


init_db()
migrate_legacy_files()

if __name__ == "__main__":
    print("\nWhatsApp Prospect Dashboard")
    print("Acesse http://127.0.0.1:5000")
    print("Use apenas contatos que autorizaram receber mensagens. O envio é manual.\n")
    try:
        webbrowser.open("http://127.0.0.1:5000")
    except Exception:
        pass
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)
